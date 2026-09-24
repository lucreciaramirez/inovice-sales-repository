from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from invoice_summary.extractor import Invoice, ParseFailure
from invoice_summary.report import (
    DETAIL_SHEET,
    FAILED_SHEET,
    GRAND_TOTAL_LABEL,
    SUMMARY_SHEET,
    build_client_summary,
    build_detail,
    export_excel,
)


def make_invoice(number, day, client, total):
    return Invoice(
        number=number,
        issue_date=day,
        client=client,
        total=Decimal(total),
        source=Path(f"{number}.pdf"),
    )


INVOICES = [
    make_invoice("00000001", date(2026, 3, 15), "Cliente Demo S.A.", "12345.67"),
    make_invoice("00000002", date(2026, 3, 18), "Ferretería Ejemplo S.R.L.", "5400.00"),
    make_invoice("00000003", date(2026, 4, 2), "Cliente Demo S.A.", "8750.50"),
    make_invoice("00000004", date(2025, 12, 20), "Kiosco Modelo", "1234.56"),
    make_invoice("00000005", date(2026, 1, 9), "Ferretería Ejemplo S.R.L.", "15000.00"),
]


def test_detail_is_sorted_by_date():
    detail = build_detail(INVOICES)
    assert list(detail["N° Factura"]) == ["0004", "0005", "0001", "0002", "0003"]


def test_client_summary_totals_and_order():
    summary = build_client_summary(INVOICES)
    clients = list(summary["Razón Social"])
    assert clients == ["Cliente Demo S.A.", "Ferretería Ejemplo S.R.L.", "Kiosco Modelo", GRAND_TOTAL_LABEL]
    assert list(summary["Importe Total"]) == [21096.17, 20400.00, 1234.56, 42730.73]
    assert list(summary["Cantidad de facturas"]) == [2, 2, 1, 5]


def test_export_excel_creates_expected_sheets(tmp_path):
    failures = [ParseFailure(Path("mala.pdf"), "Missing: total amount")]
    output = export_excel(INVOICES, failures, tmp_path / "out.xlsx")

    workbook = load_workbook(output)
    assert workbook.sheetnames == [DETAIL_SHEET, SUMMARY_SHEET, FAILED_SHEET]
    assert workbook[FAILED_SHEET]["A2"].value == "mala.pdf"


def test_export_excel_has_no_failed_sheet_when_everything_was_read(tmp_path):
    output = export_excel(INVOICES, [], tmp_path / "out.xlsx")
    assert load_workbook(output).sheetnames == [DETAIL_SHEET, SUMMARY_SHEET]


def test_export_excel_needs_invoices(tmp_path):
    with pytest.raises(ValueError):
        export_excel([], [], tmp_path / "out.xlsx")
