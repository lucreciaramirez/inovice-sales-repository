"""Build the sales report and export it to Excel."""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font

from .extractor import Invoice, ParseFailure

DETAIL_SHEET = "Detalle"
SUMMARY_SHEET = "Resumen por cliente"
FAILED_SHEET = "No procesadas"

DETAIL_COLUMNS = ["N° Factura", "Fecha de Emisión", "Razón Social", "Importe Total"]
SUMMARY_COLUMNS = ["Razón Social", "Cantidad de facturas", "Importe Total"]
GRAND_TOTAL_LABEL = "TOTAL GENERAL"


def build_detail(invoices: list[Invoice]) -> pd.DataFrame:
    """One row per invoice, ordered by date."""
    rows = [
        {
            "N° Factura": inv.short_number,
            "Fecha de Emisión": inv.issue_date,
            "Razón Social": inv.client,
            "Importe Total": float(inv.total),
        }
        for inv in sorted(invoices, key=lambda i: (i.issue_date, i.number))
    ]
    return pd.DataFrame(rows, columns=DETAIL_COLUMNS)


def build_client_summary(invoices: list[Invoice]) -> pd.DataFrame:
    """Total sales and number of invoices per client, plus a grand total row.

    Amounts are added as Decimal to avoid floating point rounding errors.
    """
    totals: dict[str, Decimal] = defaultdict(Decimal)
    counts: dict[str, int] = defaultdict(int)
    for inv in invoices:
        totals[inv.client] += inv.total
        counts[inv.client] += 1

    ranked = sorted(totals, key=lambda client: totals[client], reverse=True)
    rows = [
        {
            "Razón Social": client,
            "Cantidad de facturas": counts[client],
            "Importe Total": float(totals[client]),
        }
        for client in ranked
    ]
    rows.append(
        {
            "Razón Social": GRAND_TOTAL_LABEL,
            "Cantidad de facturas": sum(counts.values()),
            "Importe Total": float(sum(totals.values(), Decimal("0"))),
        }
    )
    return pd.DataFrame(rows, columns=SUMMARY_COLUMNS)


def build_failures(failures: list[ParseFailure]) -> pd.DataFrame:
    return pd.DataFrame(
        [{"Archivo": f.source.name, "Motivo": f.reason} for f in failures],
        columns=["Archivo", "Motivo"],
    )


def _format_sheet(ws, date_columns=(), money_columns=(), bold_last_row=False) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    headers = {cell.value: cell.column_letter for cell in ws[1]}
    for name in date_columns:
        for cell in ws[headers[name]][1:]:
            cell.number_format = "DD/MM/YYYY"
    for name in money_columns:
        for cell in ws[headers[name]][1:]:
            cell.number_format = "#,##0.00"
    if bold_last_row:
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for column in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 50)


def export_excel(
    invoices: list[Invoice], failures: list[ParseFailure], output_path: Path | str
) -> Path:
    """Write the report: a detail sheet, a per-client summary and, if any, the failed files."""
    if not invoices:
        raise ValueError("There are no invoices to export.")
    output_path = Path(output_path)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        build_detail(invoices).to_excel(writer, sheet_name=DETAIL_SHEET, index=False)
        build_client_summary(invoices).to_excel(writer, sheet_name=SUMMARY_SHEET, index=False)
        if failures:
            build_failures(failures).to_excel(writer, sheet_name=FAILED_SHEET, index=False)

        _format_sheet(
            writer.sheets[DETAIL_SHEET],
            date_columns=["Fecha de Emisión"],
            money_columns=["Importe Total"],
        )
        _format_sheet(
            writer.sheets[SUMMARY_SHEET], money_columns=["Importe Total"], bold_last_row=True
        )
        if failures:
            _format_sheet(writer.sheets[FAILED_SHEET])
    return output_path
